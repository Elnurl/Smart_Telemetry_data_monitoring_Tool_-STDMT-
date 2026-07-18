"""STDMS SOP Word template — structured procedure documents (.docx).

Matches the department SOP layout:
  STANDARD OPERATING PROCEDURE (SOP)
  - General Information
  - Process Overview
  - Process Steps (WBS / Task / Owner)
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional

SOP_TEMPLATE_ID = "stdms_sop_word_v1"

SOP_TEMPLATE_DESCRIPTION = """
STANDARD OPERATING PROCEDURE (SOP) Word template used by STDMS agent.
When creating or updating procedure/SOP files, fill these fields and emit a .docx:

General Information:
- process_title, department, contact_info, sop_id, effective_date, revision_number

Process Overview:
- process_description, purpose_scope, definitions_related

Process Steps (table):
- steps: list of {wbs, task, owner}
""".strip()

DEFAULT_SOP_FIELDS: dict[str, Any] = {
    "process_title": "",
    "department": "",
    "contact_info": "",
    "sop_id": "",
    "effective_date": "",
    "revision_number": "1",
    "process_description": "",
    "purpose_scope": "",
    "definitions_related": "",
    "steps": [],
}


def normalize_sop_fields(raw: Optional[dict] = None) -> dict[str, Any]:
    data = dict(DEFAULT_SOP_FIELDS)
    src = raw if isinstance(raw, dict) else {}
    for key in DEFAULT_SOP_FIELDS:
        if key == "steps":
            continue
        if key in src and src[key] is not None:
            data[key] = str(src[key]).strip()
    steps_in = src.get("steps") or []
    steps: list[dict[str, str]] = []
    if isinstance(steps_in, list):
        for i, item in enumerate(steps_in, start=1):
            if isinstance(item, dict):
                steps.append(
                    {
                        "wbs": str(item.get("wbs") or i).strip(),
                        "task": str(item.get("task") or "").strip(),
                        "owner": str(item.get("owner") or "").strip(),
                    }
                )
            elif isinstance(item, str) and item.strip():
                steps.append({"wbs": str(i), "task": item.strip(), "owner": ""})
    data["steps"] = steps
    if not data.get("revision_number"):
        data["revision_number"] = "1"
    return data


def suggest_sop_filename(fields: dict[str, Any]) -> str:
    sop_id = str(fields.get("sop_id") or "").strip()
    title = str(fields.get("process_title") or "").strip()
    if sop_id:
        safe = re.sub(r"[^\w.\-]+", "_", sop_id).strip("_")
        return f"{safe}.docx"
    if title:
        safe = re.sub(r"[^\w.\-]+", "_", title).strip("_")[:40]
        return f"SOP_{safe or 'procedure'}.docx"
    return "SOP_new_procedure.docx"


def build_sop_docx(fields: dict[str, Any], path: Path) -> Path:
    """Render normalized SOP fields into a Word .docx at path."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.shared import Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for SOP Word export. Install: pip install python-docx"
        ) from exc

    data = normalize_sop_fields(fields)
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run("STANDARD OPERATING PROCEDURE (SOP)")
    run.bold = True
    run.font.size = Pt(16)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _shade_cell(cell, color_hex: str = "D9D9D9") -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color_hex)
        shading.set(qn("w:val"), "clear")
        tc_pr.append(shading)

    def _section_banner(text: str) -> None:
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        cell.text = text
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(12)
        _shade_cell(cell, "BFBFBF")
        doc.add_paragraph("")

    def _kv_row(label: str, value: str) -> None:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        p.add_run(value or "—")

    # --- General Information ---
    _section_banner("General Information")
    info = doc.add_table(rows=3, cols=2)
    info.style = "Table Grid"
    pairs = [
        ("Process Title", data["process_title"]),
        ("Department", data["department"]),
        ("Contact Info", data["contact_info"]),
        ("SOP ID", data["sop_id"]),
        ("Effective Date", data["effective_date"]),
        ("Revision Number", data["revision_number"]),
    ]
    for idx, (label, value) in enumerate(pairs):
        row = info.rows[idx // 2]
        cell = row.cells[idx % 2]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(value or "—")

    doc.add_paragraph("")

    # --- Process Overview ---
    _section_banner("Process Overview")
    for label, key in (
        ("Process Description", "process_description"),
        ("Purpose & Scope", "purpose_scope"),
        ("Definitions & Related Documents", "definitions_related"),
    ):
        _kv_row(label, data[key])
        doc.add_paragraph("")

    # --- Process Steps ---
    _section_banner("Process Steps")
    steps = data["steps"] or [{"wbs": "1", "task": "", "owner": ""}]
    # pad a few empty rows like the template
    while len(steps) < 4:
        steps.append({"wbs": str(len(steps) + 1), "task": "", "owner": ""})

    steps_table = doc.add_table(rows=1 + len(steps), cols=3)
    steps_table.style = "Table Grid"
    headers = ("WBS", "Task", "Owner")
    for i, h in enumerate(headers):
        cell = steps_table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
        _shade_cell(cell, "D9D9D9")

    for r_idx, step in enumerate(steps, start=1):
        steps_table.rows[r_idx].cells[0].text = str(step.get("wbs") or r_idx)
        steps_table.rows[r_idx].cells[1].text = str(step.get("task") or "")
        steps_table.rows[r_idx].cells[2].text = str(step.get("owner") or "")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def sop_fields_to_preview_markdown(fields: dict[str, Any]) -> str:
    """Human-readable preview for draft review / chat."""
    data = normalize_sop_fields(fields)
    lines = [
        "# STANDARD OPERATING PROCEDURE (SOP)",
        "",
        "## General Information",
        f"- Process Title: {data['process_title'] or '—'}",
        f"- Department: {data['department'] or '—'}",
        f"- Contact Info: {data['contact_info'] or '—'}",
        f"- SOP ID: {data['sop_id'] or '—'}",
        f"- Effective Date: {data['effective_date'] or '—'}",
        f"- Revision Number: {data['revision_number'] or '—'}",
        "",
        "## Process Overview",
        f"**Process Description:** {data['process_description'] or '—'}",
        "",
        f"**Purpose & Scope:** {data['purpose_scope'] or '—'}",
        "",
        f"**Definitions & Related Documents:** {data['definitions_related'] or '—'}",
        "",
        "## Process Steps",
        "| WBS | Task | Owner |",
        "| --- | --- | --- |",
    ]
    for step in data["steps"] or []:
        lines.append(
            f"| {step.get('wbs') or ''} | {step.get('task') or ''} | {step.get('owner') or ''} |"
        )
    return "\n".join(lines)
