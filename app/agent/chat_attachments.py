"""Read operator-attached chat files into plain text for the agent."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger("STDMS.Agent.ChatAttachments")

# Common text / office / data types the agent can usefully summarize.
SUPPORTED_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".conf",
    ".py",
    ".js",
    ".ts",
    ".sql",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
}

MAX_FILE_BYTES = 8 * 1024 * 1024  # 8 MiB per file
MAX_CHARS_PER_FILE = 80_000
MAX_TOTAL_CHARS = 120_000
MAX_ATTACHMENTS = 8

FILE_DIALOG_FILTER = (
    "Documents & data ("
    "*.txt *.md *.markdown *.log *.csv *.tsv *.json *.jsonl "
    "*.xml *.yaml *.yml *.pdf *.docx *.py *.sql *.html"
    ");;"
    "All files (*.*)"
)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(parts).strip()


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def read_attachment(path: str | Path) -> dict[str, Any]:
    """Extract text from one file. Returns ok/error payload."""
    p = Path(path)
    name = p.name
    suffix = p.suffix.lower()
    if not p.is_file():
        return {"ok": False, "filename": name, "error": "file_not_found"}
    try:
        size = p.stat().st_size
    except OSError as exc:
        return {"ok": False, "filename": name, "error": str(exc)}
    if size > MAX_FILE_BYTES:
        return {
            "ok": False,
            "filename": name,
            "error": f"file_too_large (max {MAX_FILE_BYTES // (1024 * 1024)} MiB)",
            "size_bytes": size,
        }
    if suffix and suffix not in SUPPORTED_SUFFIXES:
        return {
            "ok": False,
            "filename": name,
            "error": f"unsupported_type ({suffix or 'no extension'})",
            "hint": "Supported: txt, md, log, csv, json, xml, yaml, pdf, docx, …",
        }

    try:
        if suffix == ".pdf":
            try:
                text = _read_pdf(p)
            except ImportError:
                return {
                    "ok": False,
                    "filename": name,
                    "error": "pypdf_not_installed",
                    "hint": "pip install pypdf",
                }
        elif suffix == ".docx":
            try:
                text = _read_docx(p)
            except ImportError:
                return {
                    "ok": False,
                    "filename": name,
                    "error": "python_docx_not_installed",
                    "hint": "pip install python-docx",
                }
        else:
            text = p.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        logger.warning("Attachment read failed for %s: %s", p, exc)
        return {"ok": False, "filename": name, "error": str(exc)}

    text = (text or "").strip()
    if not text:
        return {"ok": False, "filename": name, "error": "empty_or_unreadable"}

    truncated = len(text) > MAX_CHARS_PER_FILE
    if truncated:
        text = text[:MAX_CHARS_PER_FILE]

    return {
        "ok": True,
        "filename": name,
        "path": str(p.resolve()),
        "suffix": suffix,
        "size_bytes": size,
        "content": text,
        "truncated": truncated,
        "char_count": len(text),
    }


def build_message_with_attachments(
    user_text: str,
    attachment_paths: list[str],
) -> dict[str, Any]:
    """Combine operator text + file contents for loop.ask()."""
    question = (user_text or "").strip()
    paths = list(attachment_paths or [])[:MAX_ATTACHMENTS]
    if not question and not paths:
        return {"ok": False, "error": "empty", "message": ""}

    blocks: list[str] = []
    loaded: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    total_chars = 0

    for raw in paths:
        result = read_attachment(raw)
        if not result.get("ok"):
            errors.append(result)
            continue
        content = str(result["content"])
        remaining = MAX_TOTAL_CHARS - total_chars
        if remaining <= 0:
            errors.append(
                {
                    "ok": False,
                    "filename": result.get("filename"),
                    "error": "total_attachment_budget_exceeded",
                }
            )
            continue
        clipped = content[:remaining]
        if len(clipped) < len(content):
            result = dict(result)
            result["content"] = clipped
            result["truncated"] = True
            result["char_count"] = len(clipped)
        total_chars += len(clipped)
        loaded.append(result)
        blocks.append(
            f"### Attached file: {result['filename']}\n"
            f"(chars={result['char_count']}"
            f"{', truncated' if result.get('truncated') else ''})\n\n"
            f"{clipped}"
        )

    if not question:
        if loaded:
            question = "Please review the attached file(s) and summarize the key points."
        else:
            return {
                "ok": False,
                "error": "attachments_failed",
                "message": "",
                "errors": errors,
            }

    if blocks:
        message = (
            f"{question}\n\n"
            "--- Attached files (operator upload; use as primary context) ---\n\n"
            + "\n\n".join(blocks)
        )
    else:
        message = question

    display = question
    if loaded:
        names = ", ".join(a["filename"] for a in loaded)
        display = f"{question}\n\n[Attached: {names}]"

    return {
        "ok": True,
        "message": message,
        "display": display,
        "attachments": loaded,
        "errors": errors,
    }
